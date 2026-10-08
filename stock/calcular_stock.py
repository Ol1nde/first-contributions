#!/usr/bin/env python3
"""Calcula el stock mínimo y máximo por referencia a partir de un listado de ventas.

Reglas:
  * Periodo: los últimos 12 meses con datos (o un año natural con --año).
  * Stock mínimo = ventas de 1 mes (media mensual), redondeado hacia arriba.
  * Stock máximo = ventas de 2 meses, redondeado hacia arriba.
  * Solo se mantienen las referencias compradas por 5 clientes distintos o más
    (cliente = Código + Centro). El resto va a la hoja "Referencias a quitar".

Uso:
  python3 calcular_stock.py ventas.xlsx [-o stock.xlsx] [--año 2025]
                            [--min-clientes 5] [--meses-min 1] [--meses-max 2]
"""
import argparse
import math
import sys
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
         "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
MESES_CORTOS = [m[:3] for m in MESES]

ALIAS = {
    "año": ["ano", "año", "year", "ejercicio"],
    "cliente": ["cliente", "codigo cliente", "cod cliente", "cod. cliente", "codcliente"],
    "centro": ["centro", "cod centro", "codigo centro"],
    "nombre": ["nombre", "nombre cliente", "razon social"],
    "familia": ["familia"],
    "artículo": ["articulo", "codigo articulo", "cod articulo", "cod. articulo", "referencia", "ref"],
    "descripción": ["descripcion articulo", "descripcion", "desc articulo", "denominacion"],
    "anual": ["anual", "total", "total anual", "uds anual", "unidades"],
}

FUENTE = "Calibri Light"
AZUL_TITULO = "1F3864"
AZUL_SECCION = "3264AF"
AZUL_CABECERA = "D6E0F0"
FORMATO_NUM = '#,##0;(#,##0);"-"'
FORMATO_DEC = '#,##0.0;(#,##0.0);"-"'


def normaliza(texto):
    texto = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return " ".join(texto.lower().replace("_", " ").split())


def detecta_columnas(df):
    norm = {normaliza(c): c for c in df.columns}
    cols = {}
    for clave, alias in ALIAS.items():
        for a in alias:
            if normaliza(a) in norm:
                cols[clave] = norm[normaliza(a)]
                break
    meses = {}
    for i, (largo, corto) in enumerate(zip(MESES, MESES_CORTOS), start=1):
        for nombre in (largo, corto, "setiembre" if i == 9 else None):
            if nombre and nombre in norm:
                meses[i] = norm[nombre]
                break
    faltan = [c for c in ("cliente", "artículo") if c not in cols]
    if faltan:
        sys.exit(f"Faltan columnas obligatorias: {', '.join(faltan)}. "
                 f"Columnas encontradas: {list(df.columns)}")
    if len(meses) < 12 and "anual" not in cols:
        sys.exit("El listado necesita las 12 columnas de meses o una columna 'Anual'.")
    return cols, meses if len(meses) == 12 else {}


def lee_listado(ruta):
    ruta = Path(ruta)
    if ruta.suffix.lower() in (".csv", ".txt"):
        return pd.read_csv(ruta, sep=None, engine="python", encoding="utf-8-sig")
    hojas = pd.read_excel(ruta, sheet_name=None)
    # La hoja con más filas es la de datos brutos.
    return max(hojas.values(), key=len)


def unidades_periodo(df, cols, meses, año):
    """Devuelve (df con columna 'uds' del periodo, texto del periodo)."""
    df = df.copy()
    numerico = lambda s: pd.to_numeric(s, errors="coerce").fillna(0)

    if not meses or "año" not in cols:
        if "año" in cols:
            df[cols["año"]] = numerico(df[cols["año"]]).astype(int)
            año = año or int(df[cols["año"]].max())
            df = df[df[cols["año"]] == año]
            texto = f"Año {año}"
        else:
            texto = "Total del listado"
        df["uds"] = numerico(df[cols["anual"]])
        return df, texto

    df[cols["año"]] = numerico(df[cols["año"]]).astype(int)
    for c in meses.values():
        df[c] = numerico(df[c])

    if año:
        df = df[df[cols["año"]] == año]
        df["uds"] = df[list(meses.values())].sum(axis=1)
        return df, f"Año {año} (enero-diciembre)"

    # Últimos 12 meses: termina en el último mes con ventas del listado.
    fin = None
    for a in sorted(df[cols["año"]].unique(), reverse=True):
        sub = df[df[cols["año"]] == a]
        con_ventas = [m for m, c in meses.items() if sub[c].ne(0).any()]
        if con_ventas:
            fin = (a, max(con_ventas))
            break
    if fin is None:
        sys.exit("El listado no tiene ventas.")
    indice_fin = fin[0] * 12 + fin[1] - 1
    ventana = {(i // 12, i % 12 + 1) for i in range(indice_fin - 11, indice_fin + 1)}

    df["uds"] = 0.0
    for (a, m) in ventana:
        fila = df[cols["año"]] == a
        df.loc[fila, "uds"] += df.loc[fila, meses[m]]
    df = df[df[cols["año"]].isin({a for a, _ in ventana})]
    ini = min(ventana)
    texto = (f"Últimos 12 meses: {MESES[ini[1] - 1]} {ini[0]} - "
             f"{MESES[fin[1] - 1]} {fin[0]}")
    return df, texto


def calcula(df, cols, min_clientes, meses_min, meses_max):
    centro = df[cols["centro"]].astype(str) if "centro" in cols else ""
    df = df.assign(_art=df[cols["artículo"]].astype(str).str.strip(),
                   _cli=df[cols["cliente"]].astype(str).str.strip() + "|" + centro)

    # Neto por artículo y cliente: un cliente cuenta si compró unidades (> 0).
    por_cliente = df.groupby(["_art", "_cli"], as_index=False)["uds"].sum()
    clientes = (por_cliente[por_cliente["uds"] > 0]
                .groupby("_art")["_cli"].nunique())
    uds = por_cliente.groupby("_art")["uds"].sum()

    res = pd.DataFrame({"uds": uds})
    res["clientes"] = clientes.reindex(res.index).fillna(0).astype(int)
    for clave in ("descripción", "familia"):
        if clave in cols:
            res[clave] = (df.groupby("_art")[cols[clave]]
                          .agg(lambda s: s.dropna().astype(str).iloc[-1] if s.notna().any() else ""))
        else:
            res[clave] = ""
    res["media"] = res["uds"] / 12
    res["mínimo"] = res["media"].apply(lambda v: math.ceil(round(v * meses_min, 6)) if v > 0 else 0)
    res["máximo"] = res["media"].apply(lambda v: math.ceil(round(v * meses_max, 6)) if v > 0 else 0)

    def motivo(fila):
        if fila["uds"] <= 0:
            return "Sin ventas en el periodo"
        if fila["clientes"] < min_clientes:
            return f"Solo {fila['clientes']} cliente(s) (mínimo {min_clientes})"
        return ""

    res["motivo"] = res.apply(motivo, axis=1)
    res = res.reset_index().rename(columns={"_art": "artículo"})
    mantener = res[res["motivo"] == ""].sort_values("uds", ascending=False)
    quitar = res[res["motivo"] != ""].sort_values(["clientes", "uds"], ascending=False)
    return mantener, quitar


# ---------- Excel ----------

def estilo_hoja(ws, titulo, subtitulo, seccion, cabeceras, anchos):
    ultima = get_column_letter(len(cabeceras))
    ws["A1"] = titulo
    ws["A1"].font = Font(name=FUENTE, size=14, bold=True, color=AZUL_TITULO)
    ws["A2"] = subtitulo
    ws["A2"].font = Font(name=FUENTE, size=9, italic=True)
    ws["A3"] = seccion
    ws.merge_cells(f"A3:{ultima}3")
    ws["A3"].font = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
    ws["A3"].fill = PatternFill("solid", fgColor=AZUL_SECCION)
    for i, (cab, ancho) in enumerate(zip(cabeceras, anchos), start=1):
        c = ws.cell(row=4, column=i, value=cab)
        c.font = Font(name=FUENTE, size=9, bold=True)
        c.fill = PatternFill("solid", fgColor=AZUL_CABECERA)
        c.border = Border(bottom=Side(style="thin", color="000000"))
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{ultima}4"


def escribe_filas(ws, filas, formatos):
    cuerpo = Font(name=FUENTE, size=9)
    for r, fila in enumerate(filas, start=5):
        for c, valor in enumerate(fila, start=1):
            celda = ws.cell(row=r, column=c, value=valor)
            celda.font = cuerpo
            if formatos[c - 1]:
                celda.number_format = formatos[c - 1]
    ultima = 4 + max(len(filas), 1)
    ws.auto_filter.ref = f"A4:{get_column_letter(len(formatos))}{ultima}"


def genera_excel(mantener, quitar, salida, periodo, origen, min_clientes, meses_min, meses_max):
    wb = Workbook()

    # Hoja Stock
    ws = wb.active
    ws.title = "Stock"
    cab = ["Artículo", "Descripción", "Familia", "Nº clientes", "Uds periodo",
           "Media mensual", "Stock mínimo", "Stock máximo", "Stock actual", "A pedir", "Estado"]
    estilo_hoja(ws, "Stock mínimo y máximo por referencia",
                f"{periodo} · mínimo {meses_min} mes(es), máximo {meses_max} mes(es) de venta · "
                f"refs con {min_clientes} clientes o más",
                f"Referencias que se mantienen ({len(mantener)})", cab,
                [14, 45, 18, 10, 12, 12, 12, 12, 12, 10, 14])
    filas = []
    for i, f in enumerate(mantener.itertuples(index=False), start=5):
        filas.append([f.artículo, f.descripción, f.familia, f.clientes, f.uds, f.media,
                      f.mínimo, f.máximo, None,
                      f'=IF(I{i}="","",IF(I{i}<G{i},H{i}-I{i},0))',
                      f'=IF(I{i}="","",IF(I{i}<G{i},"Bajo mínimo",IF(I{i}>H{i},"Exceso","OK")))'])
    escribe_filas(ws, filas, [None, None, None, FORMATO_NUM, FORMATO_NUM, FORMATO_DEC,
                              FORMATO_NUM, FORMATO_NUM, FORMATO_NUM, FORMATO_NUM, None])
    entrada = PatternFill("solid", fgColor="FFF2CC")
    for r in range(5, 5 + len(filas)):
        c = ws.cell(row=r, column=9)
        c.fill = entrada
        c.font = Font(name=FUENTE, size=9, color="0000FF")
        c.alignment = Alignment(horizontal="center")
    if filas:
        rango = f"K5:K{4 + len(filas)}"
        ws.conditional_formatting.add(rango, FormulaRule(
            formula=['K5="Bajo mínimo"'], font=Font(color="9C0006"),
            fill=PatternFill("solid", fgColor="FFC7CE")))
        ws.conditional_formatting.add(rango, FormulaRule(
            formula=['K5="Exceso"'], font=Font(color="9C5700"),
            fill=PatternFill("solid", fgColor="FFEB9C")))
        ws.conditional_formatting.add(rango, FormulaRule(
            formula=['K5="OK"'], font=Font(color="006100"),
            fill=PatternFill("solid", fgColor="C6EFCE")))

    # Hoja Referencias a quitar
    wq = wb.create_sheet("Referencias a quitar")
    cab_q = ["Artículo", "Descripción", "Familia", "Nº clientes", "Uds periodo", "Motivo"]
    estilo_hoja(wq, "Referencias a quitar del stock", periodo,
                f"No cumplen los requisitos ({len(quitar)})", cab_q, [14, 45, 18, 10, 12, 34])
    escribe_filas(wq, [[f.artículo, f.descripción, f.familia, f.clientes, f.uds, f.motivo]
                       for f in quitar.itertuples(index=False)],
                  [None, None, None, FORMATO_NUM, FORMATO_NUM, None])

    # Hoja Resumen
    wr = wb.create_sheet("Resumen")
    estilo_hoja(wr, "Resumen", f"Origen: {origen}", "Criterios y resultado",
                ["Concepto", "Valor"], [42, 30])
    wr.auto_filter.ref = None
    total = len(mantener) + len(quitar)
    datos = [
        ["Periodo analizado", periodo],
        ["Stock mínimo", f"{meses_min} mes(es) de venta media"],
        ["Stock máximo", f"{meses_max} mes(es) de venta media"],
        ["Clientes mínimos por referencia", min_clientes],
        ["Referencias en el listado", total],
        ["Referencias que se mantienen", len(mantener)],
        ["Referencias a quitar", len(quitar)],
        ["  · por menos clientes del mínimo", int((quitar["uds"] > 0).sum())],
        ["  · sin ventas en el periodo", int((quitar["uds"] <= 0).sum())],
        ["Uds vendidas (refs que se mantienen)", float(mantener["uds"].sum())],
        ["Uds vendidas (refs a quitar)", float(quitar["uds"].sum())],
    ]
    escribe_filas(wr, datos, [None, FORMATO_NUM])
    for r in range(5, 5 + len(datos)):
        wr.cell(row=r, column=2).alignment = Alignment(horizontal="right")

    wb.save(salida)


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("listado", help="Listado de ventas (.xlsx, .xls o .csv)")
    p.add_argument("-o", "--salida", help="Excel de salida (por defecto: stock_<listado>.xlsx)")
    p.add_argument("--año", "--ano", dest="año", type=int, help="Año natural a usar en vez de los últimos 12 meses")
    p.add_argument("--min-clientes", type=int, default=5)
    p.add_argument("--meses-min", type=float, default=1)
    p.add_argument("--meses-max", type=float, default=2)
    a = p.parse_args()

    df = lee_listado(a.listado)
    cols, meses = detecta_columnas(df)
    df, periodo = unidades_periodo(df, cols, meses, a.año)
    mantener, quitar = calcula(df, cols, a.min_clientes, a.meses_min, a.meses_max)
    salida = a.salida or str(Path(a.listado).with_name(f"stock_{Path(a.listado).stem}.xlsx"))
    genera_excel(mantener, quitar, salida, periodo, Path(a.listado).name,
                 a.min_clientes, a.meses_min, a.meses_max)
    print(f"{periodo}\nSe mantienen: {len(mantener)} · A quitar: {len(quitar)}\nGuardado en: {salida}")


if __name__ == "__main__":
    main()
